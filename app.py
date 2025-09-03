# CORRECCIÓN: Se han eliminado caracteres invisibles (espacios de no ruptura)
# que causaban un error de sintaxis e impedían que la aplicación arrancara.
import os, io, re, base64, time
from typing import List, Tuple
from pathlib import Path

import numpy as np
import requests
from flask import Flask, request, Response
from werkzeug.utils import secure_filename

# Lectura de documentos
from pdfminer.high_level import extract_text as pdf_extract_text
from docx import Document as DocxDocument
import pypdfium2 as pdfium
from PIL import Image

# ================== Config ==================
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY", "")
OPENAI_BASE_URL = os.environ.get("OPENAI_API_URL", os.environ.get("OPENAI_BASE_URL", "https://api.openai.com/v1"))
# SOLUCIÓN: Se cambia al modelo de embeddings más compatible para evitar errores de permisos (403).
EMBED_MODEL = os.environ.get("EMBED_MODEL", "text-embedding-ada-002")
CHAT_MODEL = os.environ.get("CHAT_MODEL", "gpt-4o-mini")

# OCR opcional (si pones ENABLE_VISION_OCR=1 se fuerza siempre que no haya texto)
ENABLE_VISION_OCR = os.environ.get("ENABLE_VISION_OCR", "1") in ("1", "true", "True")
OCR_MAX_PAGES = int(os.environ.get("OCR_MAX_PAGES", "20"))
OCR_DPI = int(os.environ.get("OCR_DPI", "160"))

ALLOWED_EXT = {".pdf", ".docx", ".txt"}

# ================== Utilidades HTTP ==================
def text_response(s: str, status: int = 200) -> Response:
    return Response((s or "").strip() + "\n", status=status, mimetype="text/plain; charset=utf-8")

# ================== Utilidades de texto ==================
def allowed_file(filename: str) -> bool:
    return Path(filename.lower()).suffix in ALLOWED_EXT

def read_txt_bytes(b: bytes) -> str:
    try:
        return b.decode("utf-8", errors="ignore")
    except Exception:
        return b.decode("latin-1", errors="ignore")

def read_docx_bytes(b: bytes) -> str:
    bio = io.BytesIO(b)
    doc = DocxDocument(bio)
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join([c.text.strip() for c in row.cells]))
    return "\n".join([p for p in parts if p])

def read_pdf_bytes(b: bytes) -> str:
    bio = io.BytesIO(b)
    return pdf_extract_text(bio) or ""

def extract_text_any(filename: str, content: bytes) -> str:
    ext = Path(filename.lower()).suffix
    if ext == ".pdf":
        return read_pdf_bytes(content)
    if ext == ".docx":
        return read_docx_bytes(content)
    if ext == ".txt":
        return read_txt_bytes(content)
    return ""

def normalize_spaces(s: str) -> str:
    s = s.replace("\x00", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

def chunk_text(text: str, max_chars: int = 2500, overlap: int = 250) -> List[str]:
    paras = [p.strip() for p in re.split(r"\n\s*\n+", text) if p.strip()]
    chunks, buf, buf_len = [], [], 0
    for p in paras:
        if buf_len + len(p) + 1 > max_chars and buf:
            chunks.append("\n\n".join(buf))
            tail = chunks[-1][-overlap:] if overlap > 0 else ""
            buf = [tail, p] if tail else [p]
            buf_len = len("".join(buf))
        else:
            buf.append(p)
            buf_len += len(p) + 1
    if buf:
        chunks.append("\n\n".join(buf))
    return chunks

# ================== OpenAI wrappers ==================
def openai_embed(texts: List[str]) -> np.ndarray:
    headers = {"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"}
    payload = {"model": EMBED_MODEL, "input": texts}
    r = requests.post(f"{OPENAI_BASE_URL}/embeddings", headers=headers, json=payload, timeout=120)
    r.raise_for_status()
    data = r.json()
    vecs = [item["embedding"] for item in data["data"]]
    return np.array(vecs, dtype=np.float32)

def cosine_sim_matrix(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    a_norm = a / (np.linalg.norm(a, axis=1, keepdims=True) + 1e-8)
    b_norm = b / (np.linalg.norm(b, axis=1, keepdims=True) + 1e-8)
    return np.dot(a_norm, b_norm.T)

def build_prompt_text(instruction: str, selected_chunks: List[Tuple[int, str]]) -> List[dict]:
    corpus = "\n\n".join([f"[Fragmento {i+1}]\n{c}" for i, (_, c) in enumerate(selected_chunks)])
    system = (
        "Eres un analista experto en contratación pública del Ecuador y auditor técnico."
        "\nLee los fragmentos y cumple la instrucción con rigor y precisión."
        "\nResponde en TEXTO PLANO, sin listas, sin títulos, sin Markdown."
        "\nSi se mencionan proformas, integra comparación y conclusión de valor por dinero dentro del mismo texto."
        "\nCita entre corchetes [Fragmento #] solo cuando aporte claridad."
    )
    user = (
        f"Instrucción:\n{instruction}\n\n"
        f"Fragmentos relevantes:\n{corpus}\n\n"
        "Responde en TEXTO PLANO. No uses JSON ni listas."
    )
    return [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ]

def openai_chat(messages: List[dict]) -> str:
    headers = {"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"}
    payload = {"model": CHAT_MODEL, "messages": messages, "temperature": 0.2}
    r = requests.post(f"{OPENAI_BASE_URL}/chat/completions", headers=headers, json=payload, timeout=180)
    r.raise_for_status()
    data = r.json()
    return data["choices"][0]["message"]["content"]

# ================== OCR de respaldo (PDF imagen) ==================
def pdf_to_images(pdf_bytes: bytes, dpi: int = OCR_DPI, max_pages: int = OCR_MAX_PAGES) -> List[bytes]:
    imgs = []
    pdf = pdfium.PdfDocument(io.BytesIO(pdf_bytes))
    n = min(len(pdf), max_pages)
    for i in range(n):
        page = pdf[i]
        pil = page.render(scale=dpi/72).to_pil() # 72 dpi base
        buf = io.BytesIO()
        pil.save(buf, format="JPEG", quality=90)
        imgs.append(buf.getvalue())
    return imgs

def ocr_images_with_openai(images: List[bytes]) -> str:
    headers = {"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"}
    out = []
    batch = 4 # imágenes por request
    for i in range(0, len(images), batch):
        group = images[i:i+batch]
        content = [{"type": "text",
                    "text": "Extrae el texto legible de estas páginas en orden. Devuelve solo TEXTO PLANO, sin títulos ni listas."}]
        for img in group:
            b64 = base64.b64encode(img).decode("ascii")
            content.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})
        payload = {"model": CHAT_MODEL, "messages": [{"role": "user", "content": content}], "temperature": 0}
        r = requests.post(f"{OPENAI_BASE_URL}/chat/completions", headers=headers, json=payload, timeout=180)
        r.raise_for_status()
        out.append(r.json()["choices"][0]["message"]["content"])
    return "\n\n".join(out)

# ================== Flask App ==================
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024 # 100 MB

@app.route("/", methods=["GET"])
def health():
    return text_response("✅ EFFICON Analyzer activo.")

@app.route("/api/analyze", methods=["POST"])
def analyze():
    # LOGGING: Se añaden registros para depuración en producción
    start_time = time.time()
    print(">>> [0%] Petición recibida.")

    if not OPENAI_API_KEY:
        return text_response("OPENAI_API_KEY no configurada", 500)

    instruction = (
        (request.form.get("instruction") or "").strip()
        or (request.headers.get("X-Instruction") or "").strip()
        or (request.args.get("instruction") or "").strip()
    )
    if not instruction:
        dbg = {"content_type": request.content_type,
               "form_keys": list(request.form.keys()),
               "file_keys": list(request.files.keys())}
        return text_response(f"Falta 'instruction' (no llegó en form, header ni query). Debug: {dbg}", 400)

    upfile = request.files.get("file")
    if not upfile and "files" in request.files:
        try:
            upfile = request.files.getlist("files")[0]
        except Exception:
            upfile = None
    if not upfile or not upfile.filename:
        return text_response("Sube un archivo en el campo 'file'", 400)

    filename = secure_filename(upfile.filename)
    if not allowed_file(filename):
        return text_response(f"Extensión no permitida: {Path(filename).suffix}. Usa .pdf, .docx o .txt", 400)

    data = upfile.read()
    if not data:
        return text_response("Archivo vacío.", 400)
    
    print(f">>> [10%] Archivo '{filename}' leído en memoria.")

    try:
        text = extract_text_any(filename, data)
        text = normalize_spaces(text)
    except Exception as e:
        return text_response(f"Error extrayendo texto: {e}", 500)

    if len(text) < 50 and Path(filename).suffix.lower() == ".pdf" and ENABLE_VISION_OCR:
        try:
            pages = pdf_to_images(data, dpi=OCR_DPI, max_pages=OCR_MAX_PAGES)
            if not pages:
                return text_response("PDF sin páginas para OCR.", 422)
            ocr_text = ocr_images_with_openai(pages)
            text = normalize_spaces(ocr_text)
        except requests.HTTPError as e:
            status = e.response.status_code if e.response is not None else 502
            detail = e.response.text if e.response is not None else str(e)
            return text_response(f"OCR (visión) falló ({status}): {detail}", 502)
        except Exception as e:
            return text_response(f"OCR (visión) falló: {e}", 502)

    if len(text) < 50:
        return text_response("No se pudo extraer texto útil (¿PDF escaneado sin OCR?).", 422)

    print(f">>> [25%] Texto extraído. Longitud: {len(text)} caracteres.")
    
    full_text = f"<<{filename}>>\n{text}"
    chunks = chunk_text(full_text, max_chars=2500, overlap=250)

    try:
        # OPTIMIZACIÓN: Se combinan las dos llamadas a la API de embeddings en una sola.
        print(">>> [40%] Iniciando llamada a OpenAI Embeddings...")
        all_texts_to_embed = chunks + [instruction]
        all_vecs = openai_embed(all_texts_to_embed)
        chunk_vecs = all_vecs[:-1]
        instr_vec = all_vecs[-1:]
        print(">>> [60%] Embeddings recibidos.")

        sims = cosine_sim_matrix(instr_vec, chunk_vecs).flatten()
    except requests.HTTPError as e:
        status = e.response.status_code if e.response is not None else 502
        detail = e.response.text if e.response is not None else str(e)
        return text_response(f"Error en embeddings OpenAI ({status}): {detail}", 502)
    except Exception as e:
        return text_response(f"Error inesperado en embeddings: {e}", 502)

    K = min(10, len(chunks))
    top_idx = np.argsort(-sims)[:K]
    selected = [(int(i), chunks[int(i)]) for i in top_idx]
    
    print(f">>> [75%] Top-{K} fragmentos seleccionados.")

    messages = build_prompt_text(instruction, selected)
    try:
        print(">>> [80%] Iniciando llamada a OpenAI Chat...")
        answer = openai_chat(messages)
        print(">>> [99%] Respuesta del chat recibida.")
    except requests.HTTPError as e:
        status = e.response.status_code if e.response is not None else 502
        detail = e.response.text if e.response is not None else str(e)
        return text_response(f"Error en chat OpenAI ({status}): {detail}", 502)
    except Exception as e:
        return text_response(f"Error inesperado en chat: {e}", 502)
    
    end_time = time.time()
    total_time = end_time - start_time
    print(f">>> [100%] Proceso completado en {total_time:.2f} segundos.")
    
    return text_response(answer or "", 200)

# Este bloque asegura que app.run() solo se ejecute en desarrollo local,
# no en producción con Gunicorn, que era la causa del crash.
if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8080"))
    app.run(host="0.0.0.0", port=port)
